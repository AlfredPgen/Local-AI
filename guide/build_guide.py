r"""Build tinyGPT_learning_guide.docx from tinyGPT_learning_guide.md.

    python guide\make_figures.py      # only when the figures change
    python guide\build_guide.py       # writes ..\tinyGPT_learning_guide.docx

Needs pandoc (3.x) and python-docx. LaTeX equations become native Word
equations. Open the result in Word and accept "update fields" (or run with
--update-toc, which uses Word itself) so the table of contents shows page numbers.
"""
import argparse
import os
import shutil
import subprocess
import sys
import tempfile

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

HERE = os.path.dirname(os.path.abspath(__file__))
SOURCE = os.path.join(HERE, "tinyGPT_learning_guide.md")
OUT = os.path.join(os.path.dirname(HERE), "tinyGPT_learning_guide.docx")
PANDOC_CANDIDATES = [shutil.which("pandoc") or "", os.path.expandvars(r"%LOCALAPPDATA%\Pandoc\pandoc.exe")]
INK, ACCENT, MUTED = RGBColor(0x1F, 0x1F, 0x1F), RGBColor(0x1C, 0x5C, 0xAB), RGBColor(0x55, 0x55, 0x55)
HEADER_FILL, STRIPE_FILL, BORDER = "1C5CAB", "EEF3FA", "B7C3D0"


def pandoc_exe():
    for path in PANDOC_CANDIDATES:
        if path and os.path.isfile(path):
            return path
    raise SystemExit("pandoc not found; install it or add it to PATH")


def make_reference(pandoc, path):
    """pandoc's default reference.docx restyled: A4, Calibri, blue headings."""
    with open(path, "wb") as handle:
        handle.write(subprocess.run([pandoc, "--print-default-data-file", "reference.docx"],
                                    check=True, capture_output=True).stdout)
    doc = Document(path)
    for section in doc.sections:
        section.page_width, section.page_height = Cm(21.0), Cm(29.7)
        section.left_margin = section.right_margin = Cm(2.0)
        section.top_margin = section.bottom_margin = Cm(2.0)
    styles = doc.styles

    def font(name, size, color=INK, bold=None, italic=None, family="Calibri"):
        if name not in [s.name for s in styles]:
            return None
        st = styles[name]
        st.font.name = family
        st.font.size = Pt(size)
        st.font.color.rgb = color
        if bold is not None:
            st.font.bold = bold
        if italic is not None:
            st.font.italic = italic
        return st

    for name in ("Normal", "Body Text", "First Paragraph", "Compact"):
        st = font(name, 10.5)
        if st is not None:
            st.paragraph_format.space_after = Pt(5)
            st.paragraph_format.line_spacing = 1.1
    font("Title", 24, bold=True)
    font("Subtitle", 13, MUTED, italic=False)
    font("Date", 10.5, MUTED)
    sizes = {"Heading 1": 16, "Heading 2": 12.5, "Heading 3": 11}
    for name, size in sizes.items():
        st = font(name, size, ACCENT, bold=True)
        st.paragraph_format.space_before = Pt(16 if name == "Heading 1" else 10)
        st.paragraph_format.space_after = Pt(4)
        st.paragraph_format.keep_with_next = True
    styles["Heading 1"].paragraph_format.page_break_before = True
    font("Block Text", 10, RGBColor(0x2B, 0x3A, 0x4A))
    font("Image Caption", 9, MUTED, italic=False)
    font("Table Caption", 9, MUTED)
    font("Source Code", 8.5, family="Consolas")
    font("Verbatim Char", 9, family="Consolas")
    font("TOC Heading", 16, ACCENT, bold=True)
    doc.save(path)


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    for old in tc_pr.findall(qn("w:shd")):
        tc_pr.remove(old)
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def borders(table):
    tbl_pr = table._tbl.tblPr
    for old in tbl_pr.findall(qn("w:tblBorders")):
        tbl_pr.remove(old)
    b = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{edge}")
        e.set(qn("w:val"), "single")
        e.set(qn("w:sz"), "4")
        e.set(qn("w:space"), "0")
        e.set(qn("w:color"), BORDER)
        b.append(e)
    tbl_pr.append(b)


def style_tables(doc):
    for table in doc.tables:
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        borders(table)
        for r, row in enumerate(table.rows):
            tr_pr = row._tr.get_or_add_trPr()
            if tr_pr.find(qn("w:cantSplit")) is None:
                tr_pr.append(OxmlElement("w:cantSplit"))
            for cell in row.cells:
                if r == 0:
                    shade(cell, HEADER_FILL)
                elif r % 2 == 0:
                    shade(cell, STRIPE_FILL)
                for p in cell.paragraphs:
                    p.paragraph_format.space_after = Pt(1)
                    p.paragraph_format.space_before = Pt(1)
                    for run in p.runs:
                        run.font.size = Pt(8.5)
                        if r == 0:
                            run.font.bold = True
                            run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        if len(table.rows) > 1:  # repeat the header row on each page
            tr_pr = table.rows[0]._tr.get_or_add_trPr()
            if tr_pr.find(qn("w:tblHeader")) is None:
                tr_pr.append(OxmlElement("w:tblHeader"))


def update_toc_with_word(path):
    """Let Word compute the table of contents and page numbers (Windows only)."""
    try:
        import win32com.client  # noqa: F401
    except ImportError:
        print("pywin32 not available; open the document in Word and update fields")
        return
    import win32com.client
    word = win32com.client.DispatchEx("Word.Application")
    word.Visible = False
    word.DisplayAlerts = 0
    try:
        doc = word.Documents.Open(os.path.abspath(path))
        for toc in doc.TablesOfContents:
            toc.Update()
        doc.Save()
        doc.Close()
    finally:
        word.Quit()


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--out", default=OUT)
    ap.add_argument("--update-toc", action="store_true", help="use Word to fill in the table of contents")
    args = ap.parse_args()
    pandoc = pandoc_exe()
    with tempfile.TemporaryDirectory() as tmp:
        ref = os.path.join(tmp, "reference.docx")
        make_reference(pandoc, ref)
        raw = os.path.join(tmp, "guide.docx")
        subprocess.run([pandoc, SOURCE, "-o", raw, "--from", "markdown", "--reference-doc", ref,
                        "--number-sections", "--toc", "--toc-depth=2", "--resource-path", HERE],
                       check=True, cwd=HERE)
        doc = Document(raw)
        style_tables(doc)
        if os.path.exists(args.out):
            os.replace(args.out, args.out + ".previous")
        doc.save(args.out)
    if args.update_toc:
        update_toc_with_word(args.out)
    scrub_metadata(args.out)
    print("wrote", args.out)


def scrub_metadata(path):
    """Word writes the Office user's name into 'last modified by': clear the
    personal fields so the file can be shared."""
    doc = Document(path)
    cp = doc.core_properties
    cp.author = cp.last_modified_by = cp.comments = cp.keywords = cp.subject = cp.category = ""
    cp.revision = 1
    doc.save(path)


if __name__ == "__main__":
    sys.exit(main())
